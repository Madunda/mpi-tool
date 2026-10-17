import io
import zipfile
import re
from pathlib import Path

import pandas as pd
import streamlit as st

from core.mpi_config import MPIConfig
from core.analysis_engine import AnalysisEngine
from core.project_manager import NMPProjectManager


# ==========================================================
# PAGE CONFIGURATION
# ==========================================================

# ==========================================================
# INITIALIZATION
# ==========================================================

MPIConfig.initialize()


# ==========================================================
# HEADER
# ==========================================================

st.title("📤 Export")

st.caption(
    "Export MPI results, save portable NMPI projects, "
    "and preserve analysis outputs for reporting."
)


# ==========================================================
# SESSION DATA
# ==========================================================

df = st.session_state.get(
    "dataset"
)

result = st.session_state.get(
    "mpi_analysis_result"
)


# ==========================================================
# HELPERS
# ==========================================================

def safe_table(value):

    if value is None:
        return pd.DataFrame()

    if isinstance(
        value,
        pd.DataFrame,
    ):
        return value.copy()

    try:
        return pd.DataFrame(value)

    except Exception:
        return pd.DataFrame()


def safe_filename(name):

    cleaned = str(name).strip()

    cleaned = re.sub(
        r"[^\w\s-]",
        "",
        cleaned,
        flags=re.UNICODE,
    )

    cleaned = re.sub(
        r"\s+",
        "_",
        cleaned,
    )

    return cleaned[:100] or "nmpi_project"


def safe_sheet_name(name):

    invalid = '\\/*?:[]'

    cleaned = str(name)

    for char in invalid:
        cleaned = cleaned.replace(
            char,
            "_",
        )

    return cleaned[:31] or "Sheet"


def build_csv_zip(tables):

    output = io.BytesIO()

    with zipfile.ZipFile(
        output,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:

        for name, table in tables.items():

            if (
                table is None
                or table.empty
            ):
                continue

            filename = (
                "mpi_"
                + str(name)
                .lower()
                .replace(
                    " ",
                    "_",
                )
                .replace(
                    "-",
                    "_",
                )
                + ".csv"
            )

            archive.writestr(
                filename,
                table.to_csv(
                    index=False
                ),
            )

    output.seek(0)

    return output.getvalue()


def build_excel_workbook(
    tables,
    metadata,
):

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:

        used_sheets = set()

        for name, table in tables.items():

            if (
                table is None
                or table.empty
            ):
                continue

            sheet = safe_sheet_name(
                name
            )

            original = sheet
            counter = 1

            while sheet in used_sheets:

                suffix = f"_{counter}"

                sheet = (
                    original[
                        :31 - len(suffix)
                    ]
                    + suffix
                )

                counter += 1

            table.to_excel(
                writer,
                sheet_name=sheet,
                index=False,
            )

            used_sheets.add(
                sheet
            )

        metadata.to_excel(
            writer,
            sheet_name="Metadata",
            index=False,
        )

        # --------------------------------------------------
        # Excel formatting
        # --------------------------------------------------

        try:

            from openpyxl.styles import (
                Font,
                PatternFill,
                Alignment,
            )

            from openpyxl.utils import (
                get_column_letter,
            )

            header_fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7",
            )

            for worksheet in writer.book.worksheets:

                # Header
                for cell in worksheet[1]:

                    cell.font = Font(
                        bold=True
                    )

                    cell.fill = (
                        header_fill
                    )

                    cell.alignment = (
                        Alignment(
                            horizontal="center",
                            vertical="center",
                        )
                    )

                worksheet.freeze_panes = (
                    "A2"
                )

                if worksheet.max_row > 1:

                    worksheet.auto_filter.ref = (
                        worksheet.dimensions
                    )

                # Column width
                for column_cells in worksheet.columns:

                    max_length = 0

                    for cell in column_cells:

                        value = (
                            ""
                            if cell.value is None
                            else str(cell.value)
                        )

                        max_length = max(
                            max_length,
                            len(value),
                        )

                    letter = (
                        get_column_letter(
                            column_cells[0].column
                        )
                    )

                    worksheet.column_dimensions[
                        letter
                    ].width = min(
                        max(
                            max_length + 2,
                            10,
                        ),
                        50,
                    )

        except Exception:
            pass

    output.seek(0)

    return output.getvalue()


def build_pdf_report(
    summary_df,
    dimension_df,
    indicator_df,
    project_name,
):

    try:

        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import (
            getSampleStyleSheet,
            ParagraphStyle,
        )
        from reportlab.lib.enums import (
            TA_CENTER,
        )
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
        )

    except ImportError as error:

        raise RuntimeError(
            "PDF export requires reportlab. "
            "Install it using: pip install reportlab"
        ) from error

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "NMPI_Title",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=18,
        leading=22,
        spaceAfter=12,
    )

    subtitle_style = ParagraphStyle(
        "NMPI_Subtitle",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=9,
        leading=12,
        spaceAfter=20,
    )

    heading_style = ParagraphStyle(
        "NMPI_Heading",
        parent=styles["Heading2"],
        fontSize=12,
        leading=15,
        spaceBefore=12,
        spaceAfter=8,
    )

    body = []

    body.append(
        Paragraph(
            "National Multidimensional Poverty Index",
            title_style,
        )
    )

    body.append(
        Paragraph(
            f"Project: {project_name}",
            subtitle_style,
        )
    )

    body.append(
        Paragraph(
            "MPI Summary",
            heading_style,
        )
    )

    if not summary_df.empty:

        summary_data = [
            list(summary_df.columns)
        ]

        for _, row in summary_df.iterrows():

            summary_data.append(
                [
                    str(value)
                    for value in row.tolist()
                ]
            )

        table = Table(
            summary_data,
            repeatRows=1,
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#D9EAF7"
                        ),
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                ]
            )
        )

        body.append(table)

    if not dimension_df.empty:

        body.append(
            Paragraph(
                "Dimension Contribution",
                heading_style,
            )
        )

        display = dimension_df.copy()

        display = display[
            [
                column
                for column in [
                    "Dimension",
                    "Dimension Weight",
                    "Contribution (%)",
                ]
                if column in display.columns
            ]
        ]

        pdf_data = [
            list(display.columns)
        ]

        for _, row in display.iterrows():

            pdf_data.append(
                [
                    str(value)
                    for value in row.tolist()
                ]
            )

        table = Table(
            pdf_data,
            repeatRows=1,
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#D9EAF7"
                        ),
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                ]
            )
        )

        body.append(table)

    if not indicator_df.empty:

        body.append(
            Spacer(
                1,
                12,
            )
        )

        body.append(
            Paragraph(
                "Indicator Contribution",
                heading_style,
            )
        )

        display = indicator_df.copy()

        display = display[
            [
                column
                for column in [
                    "Dimension",
                    "Indicator",
                    "Indicator Weight",
                    "Contribution (%)",
                ]
                if column in display.columns
            ]
        ]

        pdf_data = [
            list(display.columns)
        ]

        for _, row in display.iterrows():

            pdf_data.append(
                [
                    str(value)
                    for value in row.tolist()
                ]
            )

        table = Table(
            pdf_data,
            repeatRows=1,
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#D9EAF7"
                        ),
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                ]
            )
        )

        body.append(table)

    body.append(
        Spacer(
            1,
            20,
        )
    )

    body.append(
        Paragraph(
            "NMPI Automation Tool · "
            "Alkire-Foster methodology",
            subtitle_style,
        )
    )

    document.build(body)

    output.seek(0)

    return output.getvalue()


# ==========================================================
# PROJECT SAVE / OPEN
# ==========================================================

st.divider()

st.subheader(
    "💾 Save / Open NMPI Project"
)

st.caption(
    "Save the current MPI configuration and optionally "
    "the dataset as one portable .nmpi project."
)

save_tab, open_tab = st.tabs(
    [
        "💾 Save Project",
        "📂 Open Project",
    ]
)


# ==========================================================
# SAVE PROJECT
# ==========================================================

with save_tab:

    current_config = MPIConfig.get()

    include_dataset = st.checkbox(
        "Include current dataset inside the .nmpi file",
        value=True,
        help=(
            "Including the dataset makes the project portable. "
            "For very large datasets, you may save the "
            "configuration without the dataset."
        ),
        key="include_dataset_nmpi",
    )

    project_name = current_config.get(
        "project_name",
        "New MPI Project",
    )

    filename = (
        safe_filename(
            project_name
        )
        + ".nmpi"
    )

    try:

        NMPProjectManager.validate_config(
            current_config
        )

        project_bytes = (
            NMPProjectManager.build_project_bytes(
                config=current_config,
                dataset=(
                    df
                    if include_dataset
                    else None
                ),
                dataset_name=(
                    st.session_state.get(
                        "dataset_name",
                        "dataset.csv",
                    )
                ),
            )
        )

        st.download_button(
            "⬇️ Save Current Project (.nmpi)",
            data=project_bytes,
            file_name=filename,
            mime="application/octet-stream",
            use_container_width=True,
            key="download_nmpi_project",
        )

        st.success(
            "NMPI project is ready to save."
        )

    except Exception as error:

        st.error(
            "The project cannot be saved because "
            "the current configuration is not valid."
        )

        st.warning(
            str(error)
        )


# ==========================================================
# OPEN PROJECT
# ==========================================================

with open_tab:

    uploaded_project = st.file_uploader(
        "Select an NMPI project file",
        type=["nmpi"],
        key="nmpi_project_uploader",
    )

    if uploaded_project is not None:

        try:

            project_bytes = (
                uploaded_project.getvalue()
            )

            preview = (
                NMPProjectManager
                .inspect_project_bytes(
                    project_bytes
                )
            )

            metadata = preview.get(
                "metadata",
                {},
            )

            project_name = metadata.get(
                "project_name",
                "N/A",
            )

            format_version = metadata.get(
                "format_version",
                "N/A",
            )

            dataset_included = (
                "Yes"
                if preview.get(
                    "dataset_included"
                )
                else "No"
            )

            st.info(
                f"""
                **Project:** {project_name}

                **Format version:** {format_version}

                **Dataset included:** {dataset_included}
                """
            )

            if st.button(
                "📂 Open This Project",
                type="primary",
                use_container_width=True,
                key="open_nmpi_project",
            ):

                loaded = (
                    NMPProjectManager
                    .load_project_bytes(
                        project_bytes
                    )
                )

                NMPProjectManager.validate_config(
                    loaded["config"]
                )

                MPIConfig.save(
                    loaded["config"]
                )

                loaded_dataset = (
                    loaded.get(
                        "dataset"
                    )
                )

                if loaded_dataset is not None:

                    st.session_state[
                        "dataset"
                    ] = loaded_dataset

                    st.session_state[
                        "dataset_name"
                    ] = metadata.get(
                        "dataset_name",
                        "dataset.csv",
                    )

                else:

                    st.session_state.pop(
                        "dataset",
                        None,
                    )

                    st.session_state.pop(
                        "dataset_name",
                        None,
                    )

                # ------------------------------------------
                # Invalidate old analysis results
                # ------------------------------------------

                st.session_state.pop(
                    "mpi_analysis_result",
                    None,
                )

                st.session_state.pop(
                    "mpi_disaggregation_result",
                    None,
                )

                st.session_state.pop(
                    "mpi_disaggregation",
                    None,
                )

                st.session_state.pop(
                    "mpi_dimension_contribution",
                    None,
                )

                st.session_state.pop(
                    "mpi_indicator_contribution",
                    None,
                )

                st.success(
                    "Project opened successfully. "
                    "Please run MPI Analysis again."
                )

                st.rerun()

        except Exception as error:

            st.error(
                "Could not open the selected .nmpi project."
            )

            st.warning(
                str(error)
            )


# ==========================================================
# CHECK ANALYSIS RESULTS
# ==========================================================

if df is None:

    st.divider()

    st.warning(
        "No dataset is currently loaded. "
        "Load a dataset or open an .nmpi project first."
    )

    st.stop()


if not isinstance(
    df,
    pd.DataFrame,
):

    st.error(
        "The loaded dataset is not a valid pandas DataFrame."
    )

    st.stop()


if result is None:

    st.divider()

    st.warning(
        "No MPI analysis result is available. "
        "Run MPI Analysis before exporting results."
    )

    st.stop()


# ==========================================================
# RESULT TABLES
# ==========================================================

results_df = safe_table(
    result.get(
        "data"
    )
)


try:

    summary = AnalysisEngine.summary(
        result
    )

except Exception:

    summary = {
        "Headcount (H)": result.get(
            "H"
        ),
        "Intensity (A)": result.get(
            "A"
        ),
        "MPI (M0)": result.get(
            "M0"
        ),
    }


summary_df = pd.DataFrame(
    {
        "Measure":
            list(
                summary.keys()
            ),
        "Value":
            list(
                summary.values()
            ),
    }
)


dimension_contribution = safe_table(
    result.get(
        "dimension_contribution"
    )
)


indicator_contribution = safe_table(
    result.get(
        "indicator_contribution"
    )
)


disagg_result = safe_table(
    st.session_state.get(
        "mpi_disaggregation_result"
    )
)


# ==========================================================
# AF AUDIT
# ==========================================================

try:

    audit_df = (
        AnalysisEngine.audit(
            df=df,
            result=result,
            weight_variable=result.get(
                "weight_variable"
            ),
        )
        .copy()
    )

    audit_df.insert(
        0,
        "Observation",
        audit_df.index,
    )

    audit_df.reset_index(
        drop=True,
        inplace=True,
    )

except Exception as error:

    audit_df = pd.DataFrame()

    st.warning(
        f"AF Audit could not be generated: {error}"
    )


# ==========================================================
# EXPORT TABLE COLLECTION
# ==========================================================

export_tables = {

    "MPI Summary":
        summary_df,

    "MPI Results":
        results_df,

    "Dimension Contribution":
        dimension_contribution,

    "Indicator Contribution":
        indicator_contribution,

    "Disaggregation":
        disagg_result,

    "AF Audit":
        audit_df,
}


# ==========================================================
# SUMMARY
# ==========================================================

st.divider()

st.subheader(
    "📊 Export Summary"
)

H = result.get(
    "H"
)

A = result.get(
    "A"
)

M0 = result.get(
    "M0"
)

col1, col2, col3 = st.columns(3)


with col1:

    st.metric(
        "Headcount (H)",
        (
            "N/A"
            if pd.isna(H)
            else f"{H:.4f}"
        ),
    )


with col2:

    st.metric(
        "Intensity (A)",
        (
            "N/A"
            if pd.isna(A)
            else f"{A:.4f}"
        ),
    )


with col3:

    st.metric(
        "MPI (M0)",
        (
            "N/A"
            if pd.isna(M0)
            else f"{M0:.4f}"
        ),
    )


weight_variable = result.get(
    "weight_variable"
)


if weight_variable:

    st.info(
        f"Survey weight used: `{weight_variable}`"
    )

else:

    st.warning(
        "No survey weight was selected. "
        "Equal observation weights were used."
    )


# ==========================================================
# AVAILABLE COMPONENTS
# ==========================================================

st.divider()

st.subheader(
    "📦 Available Export Components"
)

component_rows = []

for name, table in export_tables.items():

    component_rows.append(
        {
            "Component":
                name,

            "Rows":
                len(table),

            "Columns":
                len(table.columns),

            "Available":
                not table.empty,
        }
    )


st.dataframe(
    pd.DataFrame(
        component_rows
    ),
    use_container_width=True,
    hide_index=True,
)


# ==========================================================
# CSV EXPORT
# ==========================================================

st.divider()

st.subheader(
    "📄 CSV Export"
)

available_csv = [
    name
    for name, table in export_tables.items()
    if not table.empty
]


if available_csv:

    selected_csv = st.selectbox(
        "Select result to export",
        available_csv,
        key="selected_csv_export",
    )

    selected_table = (
        export_tables[
            selected_csv
        ]
    )

    csv_data = (
        selected_table
        .to_csv(
            index=False
        )
        .encode(
            "utf-8"
        )
    )

    filename = (
        "mpi_"
        + safe_filename(
            selected_csv
        )
        + ".csv"
    )

    st.download_button(
        f"⬇️ Download {selected_csv} CSV",
        data=csv_data,
        file_name=filename,
        mime="text/csv",
        use_container_width=True,
        key="download_selected_csv",
    )

else:

    st.info(
        "No result table is available for CSV export."
    )


# ==========================================================
# ZIP EXPORT
# ==========================================================

st.divider()

st.subheader(
    "📚 Complete CSV Package"
)

try:

    zip_data = build_csv_zip(
        export_tables
    )

    st.download_button(
        "⬇️ Download All MPI Results as ZIP",
        data=zip_data,
        file_name="mpi_all_results.zip",
        mime="application/zip",
        use_container_width=True,
        key="download_results_zip",
    )

except Exception as error:

    st.error(
        "Could not create ZIP export."
    )

    st.warning(
        str(error)
    )


# ==========================================================
# EXCEL EXPORT
# ==========================================================

st.divider()

st.subheader(
    "📊 Professional Excel Workbook"
)

config = MPIConfig.get()

metadata = pd.DataFrame(
    {
        "Property": [

            "Project Name",

            "Analysis Unit",

            "Poverty Cut-off",

            "Survey Weight Variable",

            "Missing Value Method",

            "Headcount (H)",

            "Intensity (A)",

            "MPI (M0)",

            "Weighted Population",

            "Weighted Poor Population",

        ],

        "Value": [

            config.get(
                "project_name",
                "N/A",
            ),

            config.get(
                "analysis_unit",
                "N/A",
            ),

            result.get(
                "cutoff"
            ),

            weight_variable
            or "None",

            result.get(
                "missing_method"
            ),

            H,

            A,

            M0,

            result.get(
                "weighted_population"
            ),

            result.get(
                "weighted_poor_population"
            ),
        ],
    }
)


try:

    excel_data = build_excel_workbook(
        export_tables,
        metadata,
    )

    st.download_button(
        "⬇️ Download Complete MPI Excel Workbook",
        data=excel_data,
        file_name="mpi_complete_results.xlsx",
        mime=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        use_container_width=True,
        key="download_excel_workbook",
    )

    st.success(
        "Professional Excel workbook is ready."
    )

except Exception as error:

    st.error(
        "Excel workbook could not be created."
    )

    st.warning(
        str(error)
    )


# ==========================================================
# PDF EXPORT
# ==========================================================

st.divider()

st.subheader(
    "📑 Professional PDF Report"
)

project_name = config.get(
    "project_name",
    "NMPI Project",
)


try:

    pdf_data = build_pdf_report(
        summary_df=summary_df,
        dimension_df=dimension_contribution,
        indicator_df=indicator_contribution,
        project_name=project_name,
    )

    st.download_button(
        "⬇️ Download MPI PDF Report",
        data=pdf_data,
        file_name=(
            safe_filename(
                project_name
            )
            + "_MPI_Report.pdf"
        ),
        mime="application/pdf",
        use_container_width=True,
        key="download_pdf_report",
    )

except Exception as error:

    st.warning(
        f"PDF export is currently unavailable: {error}"
    )


# ==========================================================
# CONTRIBUTION PREVIEW
# ==========================================================

if not dimension_contribution.empty:

    st.divider()

    st.subheader(
        "🏛️ Dimension Contribution"
    )

    st.dataframe(
        dimension_contribution,
        use_container_width=True,
        hide_index=True,
    )


if not indicator_contribution.empty:

    st.divider()

    st.subheader(
        "📌 Indicator Contribution"
    )

    st.dataframe(
        indicator_contribution,
        use_container_width=True,
        hide_index=True,
    )


# ==========================================================
# EXPORT NOTES
# ==========================================================

st.divider()

st.subheader(
    "📝 Export Contents"
)

st.markdown(
    """
    <div class="nmpi-card">

        <p>
            <strong>Excel Workbook</strong><br>
            Contains metadata, MPI summary, observation-level
            results, dimension contribution, indicator contribution,
            disaggregation and AF audit information.
        </p>

        <p>
            <strong>PDF Report</strong><br>
            Provides a compact reporting version of the main MPI
            results and contribution tables.
        </p>

        <p>
            <strong>CSV Package</strong><br>
            Provides individual machine-readable result tables.
        </p>

        <p>
            <strong>.nmpi Project</strong><br>
            Stores the MPI configuration and, when selected,
            the working dataset. Analysis results are regenerated
            after the project is opened.
        </p>

    </div>
    """,
    unsafe_allow_html=True,
)


# ==========================================================
# FOOTER
# ==========================================================

st.divider()

st.caption(
    "NMPI Automation Tool · Alkire-Foster "
    "Multidimensional Poverty Analysis"
)